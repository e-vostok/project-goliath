"""Export a names-translation spreadsheet from ``manifest.json`` (map2_3, Phase A).

Reads ``data/map/manifest.json`` and writes
``tools/map_pipeline/reference/names_table.xlsx`` with three sheets:

- ``Суша`` — LAND nodes sorted by ``name``;
- ``Суша по расположению`` — same rows sorted into a grid by ``anchor_y``
  bands of 10 units, then ``anchor_x``;
- ``Моря`` — SEA zones (their Russian names are already filled; the sheet
  is for review only).

The Project Owner fills the highlighted ``name_ru_new`` column;
``import_names.py`` (Phase B) reads the filled file back.

Run from the repository root::

    python -m tools.map_pipeline.export_names
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - exercised only without openpyxl
    Workbook = None

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA_DIR = REPO_ROOT / "data" / "map"
DEFAULT_OUT = (
    Path(__file__).resolve().parent / "reference" / "names_table.xlsx"
)

HEADERS = [
    "id",
    "key",
    "kind",
    "name",
    "name_ru_current",
    "name_ru_new",
    "comment",
    "neighbours",
    "area",
    "anchor_x",
    "anchor_y",
]

LAND_SHEET = "Суша"
GEO_SHEET = "Суша по расположению"
SEA_SHEET = "Моря"

# Width of one anchor_y band on the "Суша по расположению" sheet.
Y_BAND = 10.0
MAX_NEIGHBOURS = 4

COLUMN_WIDTHS = {
    "id": 7,
    "key": 30,
    "kind": 7,
    "name": 30,
    "name_ru_current": 30,
    "name_ru_new": 30,
    "comment": 30,
    "neighbours": 50,
    "area": 10,
    "anchor_x": 10,
    "anchor_y": 10,
}

_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_NEW_NAME_HEADER_FILL = PatternFill("solid", fgColor="FFD966")
_NEW_NAME_CELL_FILL = PatternFill("solid", fgColor="FFF2CC")


def load_manifest(data_dir: Path) -> dict:
    path = data_dir / "manifest.json"
    return json.loads(path.read_text(encoding="utf-8"))


def land_neighbour_names(manifest: dict) -> dict[int, str]:
    """``node_id -> "name1; name2; …"`` for up to 4 LAND neighbours.

    Neighbours are picked among ``land`` edges by the longest shared border
    (``len`` desc, then neighbour name); each is shown as ``name_ru`` when
    present, else ``name``.
    """
    nodes = {n["id"]: n for n in manifest["nodes"]}
    borders: dict[int, list[tuple[float, str]]] = {}
    for edge in manifest["edges"]:
        if edge["type"] != "land":
            continue
        a, b = nodes[edge["a"]], nodes[edge["b"]]
        if a["kind"] != "LAND" or b["kind"] != "LAND":
            continue
        length = edge.get("len") or 0.0
        for node, other in ((a, b), (b, a)):
            label = other["name_ru"] or other["name"]
            borders.setdefault(node["id"], []).append((length, label))
    result = {}
    for node_id, entries in borders.items():
        entries.sort(key=lambda e: (-e[0], e[1]))
        result[node_id] = "; ".join(
            label for _, label in entries[:MAX_NEIGHBOURS]
        )
    return result


def build_rows(manifest: dict) -> tuple[list[list], list[list]]:
    """Return ``(land_rows, sea_rows)`` in :data:`HEADERS` column order."""
    neighbours = land_neighbour_names(manifest)
    land_rows, sea_rows = [], []
    for node in manifest["nodes"]:
        row = [
            node["id"],
            node["key"],
            node["kind"],
            node["name"],
            node["name_ru"] or "",
            "",  # name_ru_new — filled by the Project Owner
            "",  # comment
            neighbours.get(node["id"], ""),
            round(node["area"], 2),
            round(node["anchor"][0], 2),
            round(node["anchor"][1], 2),
        ]
        (land_rows if node["kind"] == "LAND" else sea_rows).append(row)
    return land_rows, sea_rows


def _fill_sheet(ws, rows: list[list]) -> None:
    new_name_col = HEADERS.index("name_ru_new") + 1
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = (
            _NEW_NAME_HEADER_FILL if col == new_name_col else _HEADER_FILL
        )
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[
            header
        ]
    area_col = HEADERS.index("area") + 1
    for row_idx, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col == new_name_col:
                cell.fill = _NEW_NAME_CELL_FILL
            elif col == area_col:
                cell.number_format = "0.00"
            elif HEADERS[col - 1] in ("anchor_x", "anchor_y"):
                cell.number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(HEADERS))}{max(1, len(rows) + 1)}"
    )


def write_workbook(manifest: dict, out_path: Path) -> dict[str, int]:
    if Workbook is None:  # pragma: no cover
        raise SystemExit(
            "openpyxl is required: pip install -r "
            "tools/map_pipeline/requirements.txt"
        )
    land_rows, sea_rows = build_rows(manifest)
    by_geo = sorted(
        land_rows,
        key=lambda r: (
            int(r[HEADERS.index("anchor_y")] // Y_BAND),
            r[HEADERS.index("anchor_x")],
        ),
    )
    wb = Workbook()
    ws_land = wb.active
    ws_land.title = LAND_SHEET
    _fill_sheet(
        ws_land, sorted(land_rows, key=lambda r: str(r[HEADERS.index("name")]).lower())
    )
    ws_geo = wb.create_sheet(GEO_SHEET)
    _fill_sheet(ws_geo, by_geo)
    ws_sea = wb.create_sheet(SEA_SHEET)
    _fill_sheet(
        ws_sea, sorted(sea_rows, key=lambda r: str(r[HEADERS.index("key")]))
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {
        LAND_SHEET: len(land_rows),
        GEO_SHEET: len(by_geo),
        SEA_SHEET: len(sea_rows),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=DEFAULT_DATA_DIR,
        help="map data dir with manifest.json (default: data/map)",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=DEFAULT_OUT,
        help="output .xlsx path (default: reference/names_table.xlsx)",
    )
    args = parser.parse_args(argv)
    manifest = load_manifest(args.data_dir)
    counts = write_workbook(manifest, args.out)
    for sheet, count in counts.items():
        print(f"{sheet}: {count} rows")
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
