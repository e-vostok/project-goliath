"""Import filled Russian names back into ``overrides.yaml`` (map2_3, Phase B).

Reads the ``Суша`` sheet of the filled ``names_table.xlsx``: every row whose
``name_ru_new`` cell is non-empty becomes a ``names_ru`` entry keyed by the
node key. The row's ``id``/``key`` pair must match ``manifest.json`` exactly —
mismatches abort the import with every offending row listed.

Hard validation (aborts): the name must be a string, already trimmed, 1-40
characters, free of control characters and must not start or end with a
hyphen or a quote. Soft findings are reported but not fixed: Latin letters
or digits inside the name, names longer than 28 characters and duplicate
names (compared after the client search normalisation: lowercase, ё→е,
whitespace collapsed).

Existing ``names_ru`` entries are kept unless a row overwrites them — each
overwrite is listed. After the write the file is re-validated through the
``Overrides`` model and the pipeline regenerates ``manifest.json``;
``geometry.json`` and ``ids.lock.json`` must come out byte-identical
(``name_ru`` is not part of the geometry hash), otherwise the run fails.

Run from the repository root::

    python -m tools.map_pipeline.import_names tools/map_pipeline/reference/names_table.xlsx
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only without openpyxl
    load_workbook = None

from .errors import DATA_INVALID, PipelineError, PipelineFailure, print_errors
from .export_names import HEADERS, LAND_SHEET, load_manifest
from .models import load_overrides
from .pipeline import run_build

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA_DIR = REPO_ROOT / "data" / "map"
DEFAULT_OUT_DIR = Path(__file__).resolve().parent / "out"

NAME_MIN_LEN = 1
NAME_MAX_LEN = 40
# Longer names still import, but are listed for the Project Owner — the
# label budget on the map is tuned around ~28 characters.
NAME_WARN_LEN = 28

# Characters that must not open or close a name: hyphens and quotes.
_EDGE_CHARS = "-'\"`’‘“”«»"
_LATIN_OR_DIGITS_RE = re.compile(r"[A-Za-z0-9]")
_NAMES_RU_LINE_RE = re.compile(r"^names_ru\s*:")

_MAX_LISTED = 10


def _normalise_for_dupes(name: str) -> str:
    """The client-search normalisation (ё→е, lowercase, collapsed spaces)."""
    return re.sub(r"\s+", " ", name.replace("ё", "е").lower()).strip()


def read_table(
    xlsx_path: Path, land_ids: dict[str, int]
) -> tuple[dict[str, str], int, list[PipelineError], dict[str, list[str]]]:
    """Read the ``Суша`` sheet; return (names, skipped, errors, warnings).

    Only rows with a non-empty ``name_ru_new`` are used. For every used row
    the ``key`` must be a manifest LAND node and the ``id`` must equal the
    manifest id — violations are collected, not raised, so one run lists
    every offending row.
    """
    names: dict[str, str] = {}
    errors: list[PipelineError] = []
    warnings: dict[str, list[str]] = {
        "latin_or_digits": [],
        "duplicates": [],
        "over_28": [],
    }
    skipped = 0

    wb = load_workbook(xlsx_path)
    if LAND_SHEET not in wb.sheetnames:
        raise PipelineError(
            DATA_INVALID,
            f"{xlsx_path.name}: sheet {LAND_SHEET!r} not found "
            f"(sheets: {wb.sheetnames})",
        )
    ws = wb[LAND_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    col = {}
    for field in ("id", "key", "name_ru_new"):
        if field not in header:
            raise PipelineError(
                DATA_INVALID,
                f"{xlsx_path.name}: sheet {LAND_SHEET!r} has no "
                f"{field!r} column (header: {header})",
            )
        col[field] = header.index(field)

    seen_norm: dict[str, str] = {}  # normalised name -> first key
    for row_no, row in enumerate(rows, start=2):
        raw = row[col["name_ru_new"]] if col["name_ru_new"] < len(row) else None
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            skipped += 1
            continue

        key = row[col["key"]] if col["key"] < len(row) else None
        ident = row[col["id"]] if col["id"] < len(row) else None
        where = f"row {row_no} (key={key!r}, id={ident!r})"
        if not isinstance(key, str) or key not in land_ids:
            errors.append(
                PipelineError(
                    DATA_INVALID,
                    f"{LAND_SHEET} {where}: key is not a LAND node "
                    "of manifest.json",
                )
            )
            continue
        manifest_id = land_ids[key]
        try:
            row_id = int(ident)
        except (TypeError, ValueError):
            row_id = None
        if row_id != manifest_id:
            errors.append(
                PipelineError(
                    DATA_INVALID,
                    f"{LAND_SHEET} {where}: id does not match "
                    f"manifest.json (expected {manifest_id})",
                )
            )
            continue
        if key in names:
            errors.append(
                PipelineError(
                    DATA_INVALID,
                    f"{LAND_SHEET} {where}: key already named by an "
                    "earlier row",
                )
            )
            continue

        if not isinstance(raw, str):
            errors.append(
                PipelineError(
                    DATA_INVALID,
                    f"{LAND_SHEET} {where}: name_ru_new is not a "
                    f"string ({type(raw).__name__} {raw!r})",
                )
            )
            continue
        name = raw
        problems = []
        if name != name.strip():
            problems.append("not trimmed")
        if not (NAME_MIN_LEN <= len(name) <= NAME_MAX_LEN):
            problems.append(f"length {len(name)} outside 1-{NAME_MAX_LEN}")
        if any(unicodedata.category(c) in ("Cc", "Cf") for c in name):
            problems.append("contains control characters")
        if name[:1] in _EDGE_CHARS or name[-1:] in _EDGE_CHARS:
            problems.append("starts or ends with a hyphen/quote")
        if problems:
            errors.append(
                PipelineError(
                    DATA_INVALID,
                    f"{LAND_SHEET} {where}: name {name!r}: "
                    + "; ".join(problems),
                )
            )
            continue

        names[key] = name
        if _LATIN_OR_DIGITS_RE.search(name):
            warnings["latin_or_digits"].append(f"{where} name={name!r}")
        if len(name) > NAME_WARN_LEN:
            warnings["over_28"].append(
                f"{where} name={name!r} ({len(name)} chars)"
            )
        norm = _normalise_for_dupes(name)
        if norm in seen_norm:
            warnings["duplicates"].append(
                f"{where} name={name!r} duplicates key "
                f"{seen_norm[norm]!r}"
            )
        else:
            seen_norm[norm] = key

    return names, skipped, errors, warnings


def _yaml_dq(value: str) -> str:
    """Double-quoted YAML scalar — safe for «’», hyphens and «»."""
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def rewrite_names_ru(path: Path, merged: dict[str, str]) -> None:
    """Replace the ``names_ru`` block in-place, keeping all other lines.

    ``overrides.yaml`` is hand-maintained and commented, so the mapping is
    spliced in as text (sorted ``key: "name"`` lines) instead of a YAML
    round-trip that would strip comments. The file is read with CRLF
    normalisation and written back as LF bytes.
    """
    text = path.read_bytes().replace(b"\r\n", b"\n").decode("utf-8")
    lines = text.split("\n")
    start = None
    for i, line in enumerate(lines):
        if _NAMES_RU_LINE_RE.match(line):
            start = i
            break
    if start is None:
        raise PipelineError(
            DATA_INVALID,
            f"{path.name}: no top-level 'names_ru:' key found",
        )
    end = start + 1
    while end < len(lines) and lines[end][:1] in (" ", "\t"):
        end += 1
    block = ["names_ru:"] + [
        f"  {key}: {_yaml_dq(name)}" for key, name in sorted(merged.items())
    ]
    lines[start:end] = block
    path.write_bytes("\n".join(lines).encode("utf-8"))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "file",
        type=Path,
        help="filled names_table.xlsx (the «Суша» sheet is read)",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=DEFAULT_DATA_DIR,
        help="map data dir (default: data/map)",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=DEFAULT_OUT_DIR,
        help="pipeline report dir (default: tools/map_pipeline/out)",
    )
    args = parser.parse_args(argv)

    if load_workbook is None:  # pragma: no cover
        print_errors(
            [
                PipelineError(
                    DATA_INVALID,
                    "openpyxl is required: pip install -r "
                    "tools/map_pipeline/requirements.txt",
                )
            ]
        )
        return 1

    try:
        manifest = load_manifest(args.data_dir)
        land_ids = {
            n["key"]: n["id"]
            for n in manifest["nodes"]
            if n["kind"] == "LAND"
        }
        names, skipped, errors, warnings = read_table(
            args.file, land_ids
        )
        if errors:
            raise PipelineFailure(errors)

        overrides_path = args.data_dir / "overrides.yaml"
        current = load_overrides(overrides_path).names_ru
        overwritten = {
            k: (current[k], names[k])
            for k in names
            if k in current and current[k] != names[k]
        }
        merged = {**current, **names}

        # Same EOL normalisation the pipeline itself uses when diffing
        # generated files: a CRLF worktree copy is not a content change.
        def _canon(p: Path) -> bytes:
            return p.read_bytes().replace(b"\r\n", b"\n")

        geometry_path = args.data_dir / "geometry.json"
        lock_path = args.data_dir / "ids.lock.json"
        geometry_before = _canon(geometry_path)
        lock_before = _canon(lock_path)
        version_before = manifest["geometry_version"]

        rewrite_names_ru(overrides_path, merged)
        load_overrides(overrides_path)  # re-validate the edited file

        run_build(args.data_dir, args.out_dir, check=False, preview=False)

        if _canon(geometry_path) != geometry_before:
            raise PipelineError(
                DATA_INVALID,
                "geometry.json changed after a names-only edit — "
                "name_ru must not influence the geometry; investigate",
            )
        if _canon(lock_path) != lock_before:
            raise PipelineError(
                DATA_INVALID,
                "ids.lock.json changed after a names-only edit — "
                "names must not influence id assignment; investigate",
            )
        version_after = load_manifest(args.data_dir)["geometry_version"]

        print(f"imported: {len(names)} names")
        print(f"skipped (empty name_ru_new): {skipped}")
        print(f"overwritten: {len(overwritten)}")
        for key, (old, new) in sorted(overwritten.items()):
            print(f"  {key}: {old!r} -> {new!r}")
        for category, items in warnings.items():
            print(f"warnings/{category}: {len(items)}")
            for item in items[:_MAX_LISTED]:
                print(f"  {item}")
            if len(items) > _MAX_LISTED:
                print(f"  … and {len(items) - _MAX_LISTED} more")
        print(f"names_ru now holds {len(merged)} entries")
        same = "unchanged" if version_after == version_before else "CHANGED"
        print(f"geometry_version: {version_after} ({same})")
        print("manifest.json regenerated; geometry.json and "
              "ids.lock.json byte-identical")
    except PipelineFailure as failure:
        print_errors(failure.errors)
        return 1
    except PipelineError as error:
        print_errors([error])
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
