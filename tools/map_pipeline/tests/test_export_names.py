"""map2_3 Phase A: export_names produces the translation workbook."""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from tools.map_pipeline.export_names import (
    GEO_SHEET,
    HEADERS,
    LAND_SHEET,
    SEA_SHEET,
    main,
)


def _node(
    node_id: int,
    key: str,
    kind: str,
    name: str,
    anchor: list[float],
    name_ru: str | None = None,
    area: float = 4.0,
) -> dict:
    return {
        "id": node_id,
        "key": key,
        "kind": kind,
        "name": name,
        "name_ru": name_ru,
        "source_name": None if kind == "SEA" else name.replace(" ", "_"),
        "anchor": anchor,
        "bbox": [anchor[0] - 1, anchor[1] - 1, anchor[0] + 1, anchor[1] + 1],
        "area": area,
    }


@pytest.fixture
def data_dir(tmp_path):
    manifest = {
        "nodes": [
            _node(1001, "alpha", "LAND", "Alpha", [10.0, 5.0], area=4.123),
            _node(1002, "beta", "LAND", "Beta", [20.0, 15.0], name_ru="Бета"),
            _node(1003, "gamma", "LAND", "Gamma", [1.0, 8.0]),
            _node(
                2001, "sea_test", "SEA", "Тестовое море", [30.0, 40.0],
                name_ru="Тестовое море", area=10.0,
            ),
        ],
        "edges": [
            {"a": 1001, "b": 1002, "type": "land", "len": 2.0},
            {"a": 1001, "b": 1003, "type": "land", "len": 5.0},
            {"a": 1002, "b": 2001, "type": "coast", "len": 1.0},
        ],
    }
    d = tmp_path / "map"
    d.mkdir()
    (d / "manifest.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    return d


def _rows(ws):
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


def test_export_workbook(data_dir, tmp_path):
    out = tmp_path / "names_table.xlsx"
    assert main(["--data-dir", str(data_dir), "--out", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == [LAND_SHEET, GEO_SHEET, SEA_SHEET]

    land = wb[LAND_SHEET]
    rows = _rows(land)
    assert rows[0] == tuple(HEADERS)
    # "Суша" is sorted by name; name_ru_new/comment stay empty.
    assert [r[1] for r in rows[1:]] == ["alpha", "beta", "gamma"]
    assert all(r[5] is None and r[6] is None for r in rows[1:])
    assert land.freeze_panes == "A2"
    assert land.auto_filter.ref is not None
    # alpha's neighbours by longest shared border: gamma (5) then beta (2),
    # beta shown via its name_ru.
    alpha = next(r for r in rows[1:] if r[1] == "alpha")
    assert alpha[7] == "Gamma; Бета"
    assert alpha[8] == 4.12  # area rounded to 2 decimals
    assert alpha[9] == 10.0 and alpha[10] == 5.0


def test_geo_sheet_sort(data_dir, tmp_path):
    out = tmp_path / "names_table.xlsx"
    main(["--data-dir", str(data_dir), "--out", str(out)])
    rows = _rows(load_workbook(out)[GEO_SHEET])
    # Band anchor_y//10: alpha(5) and gamma(8) in band 0, beta(15) in band 1;
    # within a band rows go by anchor_x (gamma x=1 before alpha x=10).
    assert [r[1] for r in rows[1:]] == ["gamma", "alpha", "beta"]


def test_sea_sheet(data_dir, tmp_path):
    out = tmp_path / "names_table.xlsx"
    main(["--data-dir", str(data_dir), "--out", str(out)])
    rows = _rows(load_workbook(out)[SEA_SHEET])
    assert len(rows) == 2
    sea = rows[1]
    assert sea[1] == "sea_test"
    assert sea[4] == "Тестовое море"  # name_ru_current prefilled
    assert sea[5] is None  # name_ru_new stays empty for review
